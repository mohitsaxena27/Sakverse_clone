import os
from pathlib import Path

from docx import Document
from docx.oxml.ns import qn
from docx.shared import Inches
from PIL import Image
import base64
import io

def extract_chunks(docx_path):
    document = Document(docx_path)
    chunks = []

    list_styles = ['ListParagraph', 'ListBullet', 'ListNumber']

    for block in document.element.body:
        tag = block.tag.split('}')[-1]

        # Paragraph
        if tag == 'p':
            p = block
            text = ''.join([node.text for node in p.iter() if node.text])
            style = p.xpath(".//w:pStyle")
            style_id = style[0].get(qn('w:val')) if style else ''

            # Detect TOC
            if style_id and "TOC" in style_id:
                chunks.append({
                    "type": "toc",
                    "content": text.strip()
                })
            # Detect List
            elif style_id in list_styles:
                chunks.append({
                    "type": "list",
                    "content": text.strip()
                })
            # Normal Paragraph
            elif text.strip():
                chunks.append({
                    "type": "paragraph",
                    "content": text.strip()
                })

        # Table
        elif tag == 'tbl':
            table = document.tables[len([c for c in chunks if c['type'] == 'table'])]
            table_data = []
            for row in table.rows:
                table_data.append([cell.text.strip() for cell in row.cells])
            chunks.append({
                "type": "table",
                "content": "\n".join([" | ".join(row) for row in table_data])
            })

    # Images (from document.parts)
    image_idx = 0
    for rel in document.part._rels:
        rel_obj = document.part._rels[rel]
        # if "image" in rel_obj.target_ref:
        #     image_idx += 1
        #     image_data = rel_obj.target_part.blob
        #     img_b64 = base64.b64encode(image_data).decode('utf-8')
        #     chunks.append({
        #         "type": "image",
        #         "content": f"<image_{image_idx}>",
        #         "image_base64": img_b64
        #     })

    return chunks

def extract_chunks_clean(docx_path):
    document = Document(docx_path)
    chunks = []
    with open(docx_path, 'rb') as file:
        for para in document.paragraphs:
            text = para.text.strip()
            if not text:
                continue

            # Heuristic: treat list items as 'list' type
            if para.style.name.startswith("List"):
                chunk_type = "list"
            elif "TOC" in para.style.name:
                chunk_type = "toc"
            else:
                chunk_type = "paragraph"

            chunks.append({
                "type": chunk_type,
                "content": text,
                "filename":docx_path
            })

        for table in document.tables:
            table_data = []
            for row in table.rows:
                row_data = [cell.text.strip() for cell in row.cells]
                table_data.append(row_data)

            table_text = "\n".join(" | ".join(row) for row in table_data)
            chunks.append({
                "type": "table",
                "content": table_text,
                "filename": docx_path
            })
    print(chunks)
    return chunks
# Example usage
# if __name__ == "__main__":
#     docx_path = Path(__file__).resolve().parent.parent / "data" /"colt"/ "CIC User Guide_v0.3.docx"
#     # path = "./data/CIC User Guide_v0.3.docx"
#     chunks = extract_chunks_clean(docx_path)
#
#     for i, chunk in enumerate(chunks):
#         print(f"\nChunk {i+1}")
#         print(f"Type: {chunk['type']}")
#         # if chunk['type'] == 'image':
#         #     print(f"Image Base64 (first 50 chars): {chunk['image_base64']}...")
#         # else:
#         print(f"Content: {chunk['content']}")
